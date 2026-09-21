import argparse
from datetime import datetime
import sqlite3
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font
from db_helpers import BASE_DIR, DB_PATH, format_amount_nl, build_transaction_query, get_splits, calculate_balance, load_known_accounts, get_opening_balance_date_only

EXPORTS_DIR = BASE_DIR / "exports"

def parse_args():
    parser = argparse.ArgumentParser(description="Export transactions to an Excel check file.")
    parser.add_argument("--from", dest="date_from", help="Start date (YYYY-MM-DD)")
    parser.add_argument("--to", dest="date_to", help="End date (YYYY-MM-DD)")
    parser.add_argument("--category", help="Filter by 'main_type/subcategory'")
    parser.add_argument("--uncategorized", action="store_true", help="Only uncategorized transactions")
    parser.add_argument("--account", help="Filter op rekeningnummer (IBAN)")
    return parser.parse_args()

def build_query(args):
    query = """
        SELECT t.id, t.date, t.description, t.own_account, t.counter_account,
               t.code, t.direction, t.amount, t.mutation_type, t.notes,
               t.balance_after, t.category_id, c.main_type, c.subcategory
        FROM transactions t
        LEFT JOIN categories c ON t.category_id = c.id
        WHERE 1=1
    """
    params = []

    if args.date_from:
        query += " AND t.date >= ?"
        params.append(args.date_from)
    if args.date_to:
        query += " AND t.date <= ?"
        params.append(args.date_to)
    if args.uncategorized:
        query += " AND t.category_id IS NULL"
    if args.category:
        main_type, _, subcategory = args.category.partition("/")
        query += " AND c.main_type = ? AND c.subcategory = ?"
        params.append(main_type.strip())
        params.append(subcategory.strip())

    query += " ORDER BY t.date, t.id"
    return query, params

def get_splits(cursor, transaction_id):
    cursor.execute("""
        SELECT ts.amount, c.main_type, c.subcategory
        FROM transaction_splits ts
        JOIN categories c ON ts.category_id = c.id
        WHERE ts.transaction_id = ?
    """, (transaction_id,))
    return cursor.fetchall()

def get_category_lists(cursor):
    """Return (main_types, subcategories) - flat lists of distinct values."""
    cursor.execute("SELECT DISTINCT main_type FROM categories ORDER BY main_type")
    main_types = [row[0] for row in cursor.fetchall()]

    cursor.execute("SELECT DISTINCT subcategory FROM categories ORDER BY subcategory")
    subcategories = [row[0] for row in cursor.fetchall()]

    return main_types, subcategories

def export_transactions_excel(cursor, args):
    query, params = build_transaction_query(args)
    cursor.execute(query, params)
    transactions = cursor.fetchall()

    main_types, subcategories = get_category_lists(cursor)

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    filter_summary = build_filter_summary(args)
    account_filter = getattr(args, "account", None)

    current_row = 1
    ws.cell(row=current_row, column=1, value=f"Filter: {filter_summary}")
    current_row += 1

    if account_filter:
        start_balance, end_balance = get_period_bounds(args, cursor, account_filter)
        start_text = format_amount_nl(start_balance) if start_balance is not None else "onbekend"
        end_text = format_amount_nl(end_balance) if end_balance is not None else "onbekend"
        ws.cell(row=current_row, column=1,
                value=f"Saldo begin periode: {start_text}   |   Saldo einde periode: {end_text}")
        current_row += 1
    else:
        known_accounts = load_known_accounts()
        for account_number in known_accounts:
            start_balance, end_balance = get_period_bounds(args, cursor, account_number)
            if start_balance is None and end_balance is None:
                ws.cell(row=current_row, column=1,
                        value=f"{account_number} - Beginsaldo onbekend (nog niet ingesteld)")
            else:
                start_text = format_amount_nl(start_balance) if start_balance is not None else "onbekend"
                end_text = format_amount_nl(end_balance) if end_balance is not None else "onbekend"
                ws.cell(row=current_row, column=1,
                        value=f"{account_number} - Saldo begin: {start_text}   |   Saldo einde: {end_text}")
            current_row += 1

    current_row += 1  # lege regel voor leesbaarheid

    header_row = current_row
    headers = [
        "Datum", "Rekening", "Naam / Omschrijving", "Tegenrekening", "Code",
        "Af Bij", "Bedrag (EUR)", "Mutatiesoort", "Mededelingen", "Saldo na mutatie",
        "income/expense", "category"
    ]
    for col_num, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col_num, value=header)
    current_row += 1

    data_start_row = current_row

    for (tx_id, date, description, own_account, counter_account, code,
         direction, amount, mutation_type, notes, balance_after,
         category_id, main_type, subcategory) in transactions:

        af_bij = "Af" if direction == "debit" else "Bij"
        splits = get_splits(cursor, tx_id)

        rows_for_this_tx = splits if splits else [(amount, main_type, subcategory)]

        for row_amount, row_main_type, row_subcategory in rows_for_this_tx:
            ws.cell(row=current_row, column=1, value=date)
            ws.cell(row=current_row, column=2, value=own_account)
            ws.cell(row=current_row, column=3, value=description)
            ws.cell(row=current_row, column=4, value=counter_account)
            ws.cell(row=current_row, column=5, value=code)
            ws.cell(row=current_row, column=6, value=af_bij)
            ws.cell(row=current_row, column=7, value=format_amount_nl(row_amount))
            ws.cell(row=current_row, column=8, value=mutation_type)
            ws.cell(row=current_row, column=9, value=notes)
            ws.cell(row=current_row, column=10,
                    value=format_amount_nl(balance_after) if balance_after is not None else "")
            ws.cell(row=current_row, column=11, value=row_main_type or "")
            ws.cell(row=current_row, column=12, value=row_subcategory or "")
            current_row += 1

    last_data_row = current_row - 1

    # --- kleur de dropdown-kolommen lichtblauw ---
    light_blue = PatternFill(start_color="D6E9F8", end_color="D6E9F8", fill_type="solid")
    for row in range(header_row, last_data_row + 1):
        ws.cell(row=row, column=11).fill = light_blue
        ws.cell(row=row, column=12).fill = light_blue

    # --- vetgedrukte headers ---
    bold_font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        ws.cell(row=header_row, column=col).font = bold_font

    # --- autofilter ---
    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"

    # --- dropdowns (data validation) ---
    lists_ws = wb.create_sheet("_lists")
    lists_ws.sheet_state = "hidden"
    for i, value in enumerate(main_types, start=1):
        lists_ws.cell(row=i, column=1, value=value)
    for i, value in enumerate(subcategories, start=1):
        lists_ws.cell(row=i, column=2, value=value)

    main_type_range = f"_lists!$A$1:$A${len(main_types)}"
    subcategory_range = f"_lists!$B$1:$B${len(subcategories)}"

    dv_main = DataValidation(type="list", formula1=f"={main_type_range}", allow_blank=True)
    dv_sub = DataValidation(type="list", formula1=f"={subcategory_range}", allow_blank=True)
    ws.add_data_validation(dv_main)
    ws.add_data_validation(dv_sub)
    dv_main.add(f"K{header_row + 1}:K{last_data_row}")
    dv_sub.add(f"L{header_row + 1}:L{last_data_row}")

    for col_num, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_num)
        ws.column_dimensions[col_letter].width = max(12, len(header) + 2)

    EXPORTS_DIR.mkdir(exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    output_path = EXPORTS_DIR / f"export_{timestamp}.xlsx"
    wb.save(output_path)

    return output_path, last_data_row - header_row

def get_period_bounds(args, cursor, account_number):
    """Bereken saldo aan het begin en einde van de geëxporteerde periode,
    voor één rekening."""
    date_from = getattr(args, "date_from", None)
    date_to = getattr(args, "date_to", None)

    if date_from:
        # saldo op de dag vóór het begin van de periode
        day_before = (datetime.strptime(date_from, "%Y-%m-%d") - timedelta(days=1)).strftime("%Y-%m-%d")
        start_balance = calculate_balance(cursor, account_number, day_before)
    else:
        opening = get_opening_balance_date_only(cursor, account_number)
        start_balance = calculate_balance(cursor, account_number, opening) if opening else None

    end_balance = calculate_balance(cursor, account_number, date_to)

    return start_balance, end_balance


def build_filter_summary(args):
    """Leesbare samenvatting van de gebruikte filters."""
    parts = []
    date_from = getattr(args, "date_from", None)
    date_to = getattr(args, "date_to", None)
    category = getattr(args, "category", None)
    uncategorized = getattr(args, "uncategorized", False)
    account = getattr(args, "account", None)

    if account:
        parts.append(f"Rekening: {account}")
    if date_from or date_to:
        parts.append(f"Periode: {date_from or '...'} t/m {date_to or 'heden'}")
    if category:
        parts.append(f"Categorie: {category}")
    if uncategorized:
        parts.append("Alleen ongecategoriseerd")

    return " | ".join(parts) if parts else "Geen filters (alle transacties)"


if __name__ == "__main__":
    args = parse_args()
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()

    output_path, row_count = export_transactions_excel(cursor, args)
    conn.close()

    print(f"Exported {row_count} rows to {output_path}")